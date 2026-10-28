#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
排名监测：Excel 导出
======================
两种查询结果结构不同（关键词=竞品榜，ASIN=词表），所以分两个 sheet 导出，
不硬塞进一张表——列对不齐的表比没有更糟。

依赖 openpyxl（项目里 export_worker / shipment 已在用，不是新引入）。
"""

from io import BytesIO


def _pct(v, nd=1):
    if v is None or v == "":
        return ""
    try:
        return round(float(v) * 100, nd)
    except (TypeError, ValueError):
        return ""


def _num(v, nd=0):
    if v is None or v == "":
        return ""
    try:
        x = float(v)
        return int(x) if nd == 0 else round(x, nd)
    except (TypeError, ValueError):
        return ""


def _style_sheet(ws, headers, rows, widths=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="EEF1F5")
    for i, _ in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, size=10)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    # 列宽：按表头与内容的字符宽度粗估（中文按 2 个字符算）
    for i, h in enumerate(headers, 1):
        if widths and i - 1 < len(widths) and widths[i - 1]:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
            continue
        w = max([_disp_len(h)] + [_disp_len(str(r[i - 1])) for r in rows] or [0])
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 42)


def _disp_len(s: str) -> int:
    return sum(2 if ord(ch) > 127 else 1 for ch in str(s))


def export_queries(records: list) -> bytes:
    """records: [{kind, value, country, own_asin, created_at, data}] → xlsx 二进制。"""
    from openpyxl import Workbook

    kw_rows, asin_rows = [], []
    for r in records or []:
        d = r.get("data") or {}
        ts = (r.get("created_at") or "").replace("T", " ")[:16]
        if r.get("kind") == "keyword":
            my = d.get("my") or {}
            note = my.get("note") or ""
            if my.get("asin") and not my.get("share_rank"):
                note = (note or "") + (
                    f"不在前100榜；自然排名={my.get('asin_organic_rank') or '无'}"
                    if my.get("organic_checked") else "不在前100榜")
            for c in d.get("top") or []:
                kw_rows.append([
                    ts, r.get("value"), r.get("country") or "US",
                    r.get("own_asin") or "", note,
                    _num(c.get("share_rank")), c.get("asin"),
                    _num(c.get("price"), 2), _num(c.get("rating"), 1),
                    _num(c.get("review_count")), c.get("monthly_orders") or "",
                    _pct(c.get("organic_share")), _pct(c.get("sp_share")),
                    _pct(c.get("brand_share")), _pct(c.get("video_share")),
                    c.get("competition_mode") or "", _num(c.get("listed_age_days")),
                ])
        else:
            for k in d.get("keywords") or []:
                asin_rows.append([
                    ts, r.get("value"), r.get("country") or "US",
                    k.get("keyword"), _num(k.get("organic_rank")),
                    _pct(k.get("traffic_share")), _pct(k.get("natural_ratio")),
                    _num(k.get("search_volume")), k.get("health") or "",
                ])

    wb = Workbook()
    wb.remove(wb.active)     # 去掉默认空 sheet

    if kw_rows:
        ws = wb.create_sheet("关键词竞品榜")
        _style_sheet(ws, [
            "查询时间", "关键词", "站点", "我的ASIN", "我的状态", "份额排名",
            "竞品ASIN", "价格", "评分", "评论数", "月销",
            "自然%", "SP%", "品牌%", "视频%", "竞争模式", "上架天数",
        ], kw_rows)

    if asin_rows:
        ws2 = wb.create_sheet("ASIN自然排名")
        _style_sheet(ws2, [
            "查询时间", "ASIN", "站点", "关键词", "自然排名",
            "流量占比%", "自然位占比%", "搜索量", "健康度",
        ], asin_rows)

    if not kw_rows and not asin_rows:
        ws = wb.create_sheet("无数据")
        ws.append(["没有可导出的查询记录"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_single(record: dict) -> bytes:
    return export_queries([record])
